import copy
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import threading
import unittest
import urllib.request
import urllib.error
import uuid

ROOT = Path(__file__).resolve().parents[1]
# tmp/vendor/pylib is a regenerable local staging area (see tmp/stage_vendor.py);
# it comes first so a half-written vendor/pylib cannot shadow a working import.
sys.path[:0] = [str(ROOT / "tmp" / "vendor" / "pylib"), str(ROOT), str(ROOT / "vendor" / "pylib")]
import people_store
import store
import docx_gen
import xls_gen
import pdf_gen
import server
from docx import Document
from docx.oxml.ns import qn
import xlrd


def person(pid="p1", sid="20260001", name="测试甲", **extra):
    return dict(id=pid, studentId=sid, name=name, college="计算机学院", rate="100", hours="10",
                amount="800", manual=True, reason="参与科研开发与实验验证", **extra)


def make_temp_dir():
    """Sandboxed runs cannot write to the system temp dir, so stage under the repo.

    The directory is created directly instead of via tempfile.mkdtemp, because the
    0o700 ACL mkdtemp applies to its private directory is not writable here.
    """
    path = ROOT / "tmp" / "tests" / ("case-%d-%s" % (os.getpid(), uuid.uuid4().hex[:8]))
    path.mkdir(parents=True, exist_ok=True)
    return path


def drop_temp_dir(path):
    shutil.rmtree(path, ignore_errors=True)


class DependencyBootstrapTests(unittest.TestCase):
    """vendor/pylib 里残留的坏目录会让 import 成功但包是空的：

        AttributeError: module 'xlwt' has no attribute 'Workbook'

    bootstrap 必须认出这种空壳、把它排到最后，并让可用的目录生效。
    """

    def setUp(self):
        self.temp = make_temp_dir()

    def tearDown(self):
        drop_temp_dir(self.temp)

    def _run(self, script):
        """在子进程里跑，避免污染当前解释器的 sys.path / sys.modules。"""
        result = subprocess.run(
            [sys.executable, "-c", script], cwd=str(ROOT), capture_output=True,
            text=True, encoding="utf-8", errors="replace",
            env=dict(os.environ, PYTHONUTF8="1", PYTHONPATH=str(ROOT)))
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        return result.stdout

    def test_shadowing_dir_is_detected_and_healthy_dir_wins(self):
        """两种坏法都要能救回来：
        1) 目录读不了（真机上 ACL 坏掉的样子）-> 判为不健康，排到最后
        2) 目录能读但是空的（namespace package）-> 虽然判不出不健康，
           但只要可用的目录排在前面，import 就能拿到真包
        """
        root = self.temp / "app"
        (root / "vendor" / "pylib" / "xlwt").mkdir(parents=True)      # 空壳，会遮挡
        (root / "vendor" / "pylib" / "xlrd").mkdir(parents=True)
        # 用仓库附带的可用依赖目录作为健康副本；tmp/ 是运行时缓存，
        # 新克隆仓库不会包含它。
        shutil.copytree(ROOT / "vendor" / "pylib", root / "vendor" / "pylib2")

        common = (
            "import os, sys\n"
            "sys.path.insert(0, r'%s')\n"
            "import bootstrap\n"
            "bootstrap.HERE = r'%s'\n"
            "bootstrap.DEPENDENCIES = [('xlwt','xlwt'), ('xlrd','xlrd')]\n"
            "bootstrap.REQUIRED_ATTRS = {'xlwt': ('Workbook',), 'xlrd': ('open_workbook',)}\n"
            % (ROOT, root)
        )

        # 1) 空目录（能读）：靠排序救回来
        out = self._run(common + (
            "ok, note = bootstrap.ensure_dependencies(auto_install=False)\n"
            "import xlwt, xlrd\n"
            "print('结果', ok)\n"
            "print('Workbook', hasattr(xlwt, 'Workbook'))\n"
            "print('open_workbook', hasattr(xlrd, 'open_workbook'))\n"
            "print('来自', os.path.relpath(xlwt.__file__, bootstrap.HERE).replace(os.sep, '/'))\n"
        ))
        self.assertIn("结果 True", out)
        self.assertIn("Workbook True", out)
        self.assertIn("open_workbook True", out)
        self.assertIn("vendor/pylib2/xlwt", out)

        # 2) 目录读不了：必须被判为不健康，并排到可用目录后面
        out = self._run(common + (
            "bad = os.path.join(bootstrap.HERE, 'vendor', 'pylib')\n"
            "real = bootstrap._healthy_dir\n"
            "bootstrap._healthy_dir = lambda p: False if os.path.abspath(p) == os.path.abspath(bad) else real(p)\n"
            "print('健康判定', bootstrap._healthy_dir(bad))\n"
            "ok, note = bootstrap.ensure_dependencies(auto_install=False)\n"
            "import xlwt\n"
            "print('结果', ok)\n"
            "print('Workbook', hasattr(xlwt, 'Workbook'))\n"
            "print('来自', os.path.relpath(xlwt.__file__, bootstrap.HERE).replace(os.sep, '/'))\n"
        ))
        self.assertIn("健康判定 False", out)
        self.assertIn("结果 True", out)
        self.assertIn("Workbook True", out)
        self.assertIn("vendor/pylib2/xlwt", out)

    def test_primary_vendor_dir_wins_and_reinstall_avoids_it(self):
        """有多个可用目录时按 vendor/pylib -> vendor/pylib2 的顺序取；
        但重装要往 vendor/pylib2 装，免得再往坏目录里写。"""
        root = self.temp / "app"
        first = root / "vendor" / "pylib"
        second = root / "vendor" / "pylib2"
        for path in (first, second):
            (path / "xlwt").mkdir(parents=True)
            (path / "xlwt" / "__init__.py").write_text("", encoding="utf-8")
        (first / "xlwt" / "__init__.py").write_text('MARK = "PRIMARY"\n', encoding="utf-8")
        (second / "xlwt" / "__init__.py").write_text('MARK = "SECOND"\n', encoding="utf-8")

        script = (
            "import os, sys\n"
            "sys.path.insert(0, r'%s')\n"
            "import bootstrap\n"
            "bootstrap.HERE = r'%s'\n"
            "bootstrap.DEPENDENCIES = [('xlwt','xlwt')]\n"
            "bootstrap.REQUIRED_ATTRS = {'xlwt': ()}\n"
            "bootstrap.ensure_dependencies(auto_install=False)\n"
            "import xlwt\n"
            "print('用的', xlwt.MARK)\n"
            "print('重装目标', os.path.relpath(bootstrap.install_target(), bootstrap.HERE))\n"
            % (ROOT, root)
        )
        out = self._run(script)
        self.assertIn("用的 PRIMARY", out)
        self.assertIn("pylib2", out.replace("\\", "/"))

    def test_root_vendor_dir_does_not_shadow_installed_packages(self):
        """vendor/ 自己也是个目录：它排在 sys.path 上不能挡住已装的包。"""
        script = (
            "import sys\n"
            "sys.path.insert(0, r'%s')\n"
            "import bootstrap\n"
            "bootstrap.ensure_dependencies(auto_install=False)\n"
            "import docx, reportlab\n"
            "print('docx', bool(docx.Document))\n"
            "print('reportlab', bool(reportlab.__version__))\n"
            % ROOT
        )
        out = self._run(script)
        self.assertIn("docx True", out)
        self.assertIn("reportlab True", out)


class PersistentLibraryTests(unittest.TestCase):
    def setUp(self):
        self.temp = make_temp_dir()
        self.path = str(self.temp / "people.sqlite3")
        self.legacy = {"settings": {"period": "2026年9月"}, "students": [person()], "step": 2}
        people_store.initialize(self.path, self.legacy)

    def tearDown(self):
        drop_temp_dir(self.temp)

    def mutate(self, action, **body):
        return people_store.mutate(self.path, action, dict(body, revision=people_store.load(self.path)["revision"]))

    def test_migration_once_and_sequential_additions_survive_reopen(self):
        self.mutate("person.save", person=person("p2", "", "同名"))
        self.mutate("person.save", person=person("p3", "", "同名"))
        people_store.initialize(self.path, {"students": []})
        state = people_store.load(self.path)
        self.assertEqual([p["id"] for p in state["library"]["people"]], ["p1", "p2", "p3"])
        self.mutate("person.save", person={"id": "p2", "name": "改名后", "customFields": {"备注": "长期"}})
        self.assertEqual(people_store.load(self.path)["library"]["people"][1]["name"], "改名后")

    def test_empty_workspace_does_not_delete_people(self):
        self.mutate("workspace", workspace={"students": [], "settings": {"period": "2026年10月"}})
        state = people_store.load(self.path)
        self.assertEqual(state["students"], [])
        self.assertEqual(len(state["library"]["people"]), 1)
        self.mutate("workspace", workspace={"students": [dict(state["library"]["people"][0], checked=True)]})
        self.assertEqual(people_store.load(self.path)["students"][0]["id"], "p1")

    def test_identity_is_exclusive_and_project_groups_overlap(self):
        """身份组互斥（防止重复发放），项目组可以和身份组、其他项目组重叠。"""
        state = self.mutate("group.save", name="国家重点研发", kind="project")
        groups = {g["name"]: g["id"] for g in state["library"]["groups"]}
        grad, undergrad, project = groups["研究生"], groups["本科生"], groups["国家重点研发"]
        self.assertEqual([g["kind"] for g in state["library"]["groups"] if g["name"] == "国家重点研发"], ["project"])

        state = self.mutate("person.save", person={"id": "p1", "identityId": grad, "projectIds": [project]})
        saved = state["library"]["people"][0]
        self.assertEqual((saved["identityId"], saved["projectIds"]), (grad, [project]))

        # 改动别的字段不能把分组弄丢
        state = self.mutate("person.save", person={"id": "p1", "amount": "1200"})
        saved = state["library"]["people"][0]
        self.assertEqual((saved["identityId"], saved["projectIds"], saved["amount"]), (grad, [project], "1200"))

        # 换成另一个身份组：仍然只有一个身份
        state = self.mutate("person.save", person={"id": "p1", "identityId": undergrad})
        self.assertEqual(state["library"]["people"][0]["identityId"], undergrad)

        # 项目组不能当身份组用
        with self.assertRaises(ValueError):
            self.mutate("person.save", person={"id": "p1", "identityId": project})

    def test_group_delete_keeps_people_and_clears_membership(self):
        state = self.mutate("group.save", name="横向课题A", kind="project")
        groups = {g["name"]: g["id"] for g in state["library"]["groups"]}
        grad, project = groups["研究生"], groups["横向课题A"]
        self.mutate("person.save", person={"id": "p1", "identityId": grad, "projectIds": [project]})
        state = self.mutate("group.delete", id=project)
        saved = state["library"]["people"][0]
        self.assertEqual(saved["projectIds"], [])
        self.assertEqual(saved["identityId"], grad)
        state = self.mutate("group.delete", id=grad)
        saved = state["library"]["people"][0]
        self.assertEqual(saved["identityId"], "")
        self.assertEqual(len(state["library"]["people"]), 1)

    def test_group_roster_import_is_batch_and_blocks_cross_group_duplicates(self):
        state = people_store.load(self.path)
        groups = {g["name"]: g["id"] for g in state["library"]["groups"]}
        grad, undergrad = groups["研究生"], groups["本科生"]

        state = self.mutate("group.roster", id=grad, kind="identity",
                            people=[person("p1", "20260001", "甲"), person("p2", "20260002", "乙")])
        self.assertEqual([p["identityId"] for p in state["library"]["people"]], [grad, grad])

        # 同一个人不能同时出现在两个身份组里：整批拦下来，不产生半截数据
        with self.assertRaises(ValueError):
            self.mutate("group.roster", id=undergrad, kind="identity",
                        people=[person("p3", "20260003", "丙"), person("p4", "20260001", "甲")])
        self.assertEqual(len(people_store.load(self.path)["library"]["people"]), 2)

        state = self.mutate("group.roster", id=undergrad, kind="identity",
                            people=[person("p3", "20260003", "丙")])
        self.assertEqual(len(state["library"]["people"]), 3)

        # 替换式导入：本组没再出现的人保留档案，只是脱离该组
        state = self.mutate("group.roster", id=grad, kind="identity", replace=True,
                            people=[person("p1", "20260001", "甲")])
        self.assertEqual(len(state["library"]["people"]), 3)
        moved = [p for p in state["library"]["people"] if p["name"] == "乙"][0]
        self.assertEqual(moved["identityId"], "")

    def test_delete_is_explicit_stale_saves_cannot_resurrect(self):
        self.mutate("person.save", person=person("p2", "20260002", "测试乙"))
        stale = people_store.load(self.path)
        state = self.mutate("person.delete", id="p1")
        self.assertEqual([p["id"] for p in state["library"]["people"]], ["p2"])
        self.assertEqual(state["students"], [])
        with self.assertRaises(people_store.ConflictError):
            people_store.mutate(self.path, "workspace", {"revision": stale["revision"], "workspace": stale})
        with self.assertRaises(people_store.ConflictError):
            self.mutate("person.save", person=person())

    def test_failed_transaction_preserves_existing_people_and_revision(self):
        old = people_store.load(self.path)
        bad = person("p2", "20260002", "测试乙", identityId="missing")
        with self.assertRaises(ValueError):
            self.mutate("workspace", workspace={"students": [person(name="不应保存"), bad]})
        self.assertEqual(people_store.load(self.path), old)

    def test_student_id_identity_and_backup_roundtrip(self):
        state = self.mutate("person.save", person=person("new-id", "20260001", "更新姓名"))
        self.assertEqual(len(state["library"]["people"]), 1)
        self.assertEqual(state["library"]["people"][0]["id"], "p1")
        groups = {g["name"]: g["id"] for g in state["library"]["groups"]}
        self.mutate("group.save", name="横向课题A", kind="project")
        project = [g["id"] for g in people_store.load(self.path)["library"]["groups"] if g["name"] == "横向课题A"][0]
        self.mutate("person.save", person={"id": "p1", "identityId": groups["研究生"], "projectIds": [project]})
        backup = people_store.backup(self.path)
        other = str(self.temp / "restored.sqlite3")
        people_store.initialize(other, {"students": []})
        restored = people_store.mutate(other, "library.import", {"revision": 0, "backup": backup})
        saved = restored["library"]["people"][0]
        self.assertEqual(saved["name"], "更新姓名")
        restored_groups = {g["name"]: g["id"] for g in restored["library"]["groups"]}
        self.assertEqual(saved["identityId"], restored_groups["研究生"])
        self.assertEqual(saved["projectIds"], [restored_groups["横向课题A"]])


class ApiTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp = make_temp_dir()
        cls.old_paths = (store.APP_DIR, store.SETTINGS_PATH, store.DATABASE_PATH, store.OUTPUT_DIR, server.OUTPUT_DIR)
        store.APP_DIR = str(cls.temp)
        store.SETTINGS_PATH = str(cls.temp / "settings.json")
        store.DATABASE_PATH = str(cls.temp / "people.sqlite3")
        store.OUTPUT_DIR = server.OUTPUT_DIR = str(cls.temp / "output")
        cls.http = server.ThreadingHTTPServer(("127.0.0.1", 0), server.Handler)
        cls.thread = threading.Thread(target=cls.http.serve_forever, daemon=True)
        cls.thread.start()
        cls.url = "http://127.0.0.1:%d" % cls.http.server_port

    @classmethod
    def tearDownClass(cls):
        cls.http.shutdown()
        cls.http.server_close()
        cls.thread.join()
        store.APP_DIR, store.SETTINGS_PATH, store.DATABASE_PATH, store.OUTPUT_DIR, server.OUTPUT_DIR = cls.old_paths
        drop_temp_dir(cls.temp)

    def request(self, path, body=None):
        data = json.dumps(body, ensure_ascii=False).encode("utf-8") if body is not None else None
        req = urllib.request.Request(self.url + path, data=data, headers={"Content-Type": "application/json"})
        with urllib.request.urlopen(req) as response:
            return json.load(response)

    def test_api_persistence_and_reset_keeps_library(self):
        state = self.request("/api/settings")
        response = self.request("/api/people/save", {"revision": state["revision"], "person": person("api-person")})
        state = response["workspace"]
        state["students"] = [person("api-person")]
        saved = self.request("/api/settings", state)["settings"]
        self.assertEqual(saved["students"][0]["name"], "测试甲")
        with self.assertRaises(urllib.error.HTTPError) as error:
            self.request("/api/settings", state)
        self.assertEqual(error.exception.code, 409)
        reset = self.request("/api/settings/reset", {"revision": saved["revision"]})["workspace"]
        self.assertEqual(reset["students"], [])
        self.assertTrue(any(p["name"] == "测试甲" for p in reset["library"]["people"]))

    def test_blank_student_ids_batch_fill_and_exact_total(self):
        students = [person("a", "", "甲"), person("b", "", "乙"), person("c", "", "丙")]
        preview = self.request("/api/batch/preview", {"field": "amount", "students": students, "values": [11,22,33]})
        self.assertEqual([p["amount"] for p in preview["items"]], ["11", "22", "33"])
        preview = self.request("/api/batch/preview", {"field": "amount", "students": students, "total": "100"})
        self.assertEqual([p["amount"] for p in preview["items"]], ["33.34", "33.33", "33.33"])
        self.assertEqual(round(sum(float(p["amount"]) for p in preview["items"]), 2), 100)

    def test_export_both_templates_and_exact_submission_columns(self):
        students = [person("z", "20260003", "测试丙"), person("y", "20260002", "测试乙", checked=False), person("x", "20260001", "测试甲")]
        for kind in ("research", "non_research"):
            with self.subTest(kind=kind):
                payload = {"settings": {"projectType":kind,"period":"2026年9月", "unitName":"测试单位", "projectName":"测试项目"},
                           "students":students, "kinds":["docx","pdf","xls","submission"]}
                results = self.request("/api/export", payload)["results"]
                doc = Document(results["docx"]["path"])
                layout = docx_gen.inspect_template(doc)
                self.assertEqual(layout["research"], kind == "research")
                text = "".join(n.text or "" for n in doc.element.iter(qn("w:t")))
                self.assertIn("测试甲", text)
                self.assertNotIn("测试乙", text)
                # 金额带千分位，和页面上看到的一致
                self.assertIn("1,600", text)
                self.assertLess(text.index("测试甲"), text.index("测试丙"))
                self.assertEqual(Path(results["pdf"]["path"]).read_bytes()[:5], b"%PDF-")
                sheet = xlrd.open_workbook(results["submission"]["path"]).sheet_by_index(0)
                self.assertEqual(sheet.row_values(0), ["学号", "姓名", "助研津贴(三兼费)"])
                self.assertEqual((sheet.nrows,sheet.ncols), (3,3))
                self.assertEqual(sheet.cell_value(1,0), "20260001")
                roster = xlrd.open_workbook(results["xls"]["path"]).sheet_by_index(0)
                self.assertEqual((roster.nrows,roster.ncols), (4,7))
                self.assertEqual(roster.cell_value(2,0), "")

    def test_reason_template_never_names_people_and_marks_todo(self):
        """事由按模板生成：不写"xxx等N位"，待填的地方加高亮。"""
        settings = {"period": "2026年9月", "projectName": "测试项目"}
        student = {"name": "测试甲", "studentId": "20260001", "identityName": "本科生"}

        # 默认模板：身份 + 姓名 + 两处待填写占位
        text = docx_gen.build_reason(student, settings, project_name="")
        self.assertEqual(text, "本科生测试甲参与了【待填写项目名称】项目，完成了【待填写工作内容】工作。")
        self.assertNotIn("等", text)

        # 调用方给了项目名（按项目组发放）：直接写进去
        text = docx_gen.build_reason(student, settings, project_name="国家重点研发")
        self.assertIn("参与了国家重点研发项目", text)

        # 填了工作内容就不再高亮
        filled = dict(settings, workContent="数据清洗与模型训练")
        self.assertIn("完成了数据清洗与模型训练工作", docx_gen.build_reason(student, filled))

        # 本人自己写的优先
        own = dict(student, reason="自己写的事由")
        self.assertEqual(docx_gen.build_reason(own, settings), "自己写的事由")

        # 说明段落不再点名
        note = docx_gen.build_note(settings, [student])
        self.assertNotIn("测试甲", note)
        self.assertNotIn("等", note)
        self.assertIn("见下表", note)

    def test_export_marks_todo_highlight_in_word_and_pdf(self):
        # reason 留空才会走模板；自己填了事由的以自己写的为准（见下面第二个断言）
        blank = person("h1", "20260001", "测试甲", identityName="本科生")
        blank["reason"] = ""
        students = [blank]
        payload = {"settings": {"projectType": "research", "period": "2026年9月"},
                   "students": students, "kinds": ["docx", "pdf"]}
        results = self.request("/api/export", payload)["results"]

        doc = Document(results["docx"]["path"])
        highlighted, plain = [], []
        for run in doc.element.iter(qn("w:r")):
            body = "".join(n.text or "" for n in run.iter(qn("w:t")))
            if not body:
                continue
            rPr = run.find(qn("w:rPr"))
            mark = rPr.find(qn("w:highlight")) if rPr is not None else None
            (highlighted if mark is not None else plain).append(body)
        joined = "".join(plain + highlighted)
        # 方括号只作为输入标记，导出时要去掉
        self.assertNotIn("【", joined)
        self.assertNotIn("】", joined)
        self.assertIn("待填写项目名称", highlighted)
        self.assertIn("待填写工作内容", highlighted)
        self.assertIn("本科生测试甲参与了", "".join(plain))

        # 自己写了事由的：原样使用，不加高亮
        own = [person("h2", "20260002", "测试乙", identityName="研究生")]
        payload = {"settings": {"projectType": "research", "period": "2026年9月"},
                   "students": own, "kinds": ["docx"]}
        path = self.request("/api/export", payload)["results"]["docx"]["path"]
        text = "".join(n.text or "" for n in Document(path).element.iter(qn("w:t")))
        self.assertIn("参与科研开发与实验验证", text)
        self.assertNotIn("待填写", text)

    def test_export_project_group_supplies_project_name_and_leaves_work_blank(self):
        """有项目组时，项目名直接写进事由，工作内容留空高亮。"""
        state = self.request("/api/settings")
        groups = {g["name"]: g["id"] for g in state["library"]["groups"]}
        state = self.request("/api/groups/save", {
            "revision": state["revision"], "name": "xx项目组", "kind": "project"})["workspace"]
        project = [g["id"] for g in state["library"]["groups"] if g["name"] == "xx项目组"][0]

        students = [person("h1", "20260001", "测试甲", identityName="研究生")]
        students[0]["reason"] = ""
        base = {"projectType": "research", "period": "2026年9月"}

        # 套用项目组：项目名写进去，工作内容留高亮
        payload = {"settings": dict(base, reasonProjectGroup=project), "students": students, "kinds": ["docx"]}
        path = self.request("/api/export", payload)["results"]["docx"]["path"]
        text = "".join(n.text or "" for n in Document(path).element.iter(qn("w:t")))
        self.assertIn("参与了xx项目组项目", text)
        self.assertIn("待填写工作内容", text)
        self.assertNotIn("待填写项目名称", text)

        # 不套用（界面上「不套用项目组」的哨兵值）：项目名也留高亮
        payload = {"settings": dict(base, reasonProjectGroup="__off__"), "students": students, "kinds": ["docx"]}
        path = self.request("/api/export", payload)["results"]["docx"]["path"]
        text = "".join(n.text or "" for n in Document(path).element.iter(qn("w:t")))
        self.assertIn("待填写项目名称", text)
        self.assertNotIn("xx项目组", text)

        # 填了工作内容就不再高亮
        payload = {"settings": dict(base, reasonProjectGroup=project, workContent="数据清洗"),
                   "students": students, "kinds": ["docx"]}
        path = self.request("/api/export", payload)["results"]["docx"]["path"]
        text = "".join(n.text or "" for n in Document(path).element.iter(qn("w:t")))
        self.assertIn("完成了数据清洗工作", text)
        self.assertNotIn("待填写", text)

    def test_roster_is_ordered_by_student_id(self):
        """名单按学号升序：明细表、名单表、上传表三处顺序要一致。"""
        students = [person("c", "20260003", "丙"), person("a", "20260001", "甲"),
                    person("b", "20260002", "乙")]
        payload = {"settings": {"projectType": "research", "period": "2026年9月"},
                   "students": students, "kinds": ["docx", "xls", "submission"]}
        results = self.request("/api/export", payload)["results"]

        doc = Document(results["docx"]["path"])
        text = "".join(n.text or "" for n in doc.element.iter(qn("w:t")))
        self.assertLess(text.index("甲"), text.index("乙"))
        self.assertLess(text.index("乙"), text.index("丙"))

        roster = xlrd.open_workbook(results["xls"]["path"]).sheet_by_index(0)
        names = [roster.cell_value(row, 2) for row in range(1, roster.nrows)]
        self.assertEqual(names, ["甲", "乙", "丙"])

        submit = xlrd.open_workbook(results["submission"]["path"]).sheet_by_index(0)
        ids = [submit.cell_value(row, 0) for row in range(1, submit.nrows)]
        self.assertEqual(ids, ["20260001", "20260002", "20260003"])

    def test_detail_rows_are_split_and_reason_column_is_widened(self):
        """科研模板的明细行是纵向合并的（两人挤一格）+ 事由列过窄，导出前都要处理：
        - 拆开纵向合并，一人一格、事由各归各行
        - 事由列按最长那句话加宽
        """
        # 事由留空才会用模板生成（自己填了事由的以自己写的为准，那种情况不必加宽）
        students = []
        for pid, sid, name in (("a", "20260001", "甲"), ("b", "20260002", "乙"), ("c", "20260003", "丙")):
            item = person(pid, sid, name)
            item["reason"] = ""
            students.append(item)
        payload = {"settings": {"projectType": "research", "period": "2026年9月",
                                "workContent": "数据清洗与模型训练"},
                   "students": students, "kinds": ["docx"], "projectName": "国家重点研发"}
        path = self.request("/api/export", payload)["results"]["docx"]["path"]
        doc = Document(path)
        table = doc.tables[0]._tbl
        layout = docx_gen.inspect_template(doc)
        rows = table.findall(qn("w:tr"))

        # 事由列加宽了（模板原始 2705 twips）
        grid = table.find(qn("w:tblGrid"))
        widths = [int(c.get(qn("w:w"))) for c in grid]
        reason_w = sum(widths[layout["reason_col"]:layout["reason_col"] + layout["reason_span"]])
        self.assertGreater(reason_w, 2705, "事由列没有加宽：%d" % reason_w)
        self.assertEqual(sum(widths), sum(docx_gen.inspect_template(
            Document(store.current_docx_template("research")))["widths"]), "总宽不该变")

        # 明细行里不能再有纵向合并，且每人一行、事由各归各行
        detail = []
        for row in rows[layout["first"]:]:
            cells = row.findall(qn("w:tc"))
            if len(cells) != 5:
                continue           # 合计行/说明行，不是明细
            name = "".join(n.text or "" for n in cells[1].iter(qn("w:t")))
            if not name:
                continue
            for cell in cells:
                tcPr = cell.find(qn("w:tcPr"))
                self.assertIsNone(tcPr.find(qn("w:vMerge")) if tcPr is not None else None,
                                  "明细行还有纵向合并的格子")
            detail.append((name, "".join(n.text or "" for n in cells[4].iter(qn("w:t")))))

        self.assertEqual(len(detail), len(students), detail)
        for name, reason in detail:
            self.assertIn(name + "参与了", reason)
            self.assertIn("国家重点研发", reason)
        self.assertEqual(len(set(reason for _, reason in detail)), len(students),
                         "两个人的事由不该是同一格内容")

    def test_export_rejects_missing_name_and_missing_submission_id(self):
        for students in ([person(name="")], [person(sid="")]):
            with self.assertRaises(urllib.error.HTTPError) as error:
                self.request("/api/export", {"students":students,"kinds":["submission"]})
            self.assertEqual(error.exception.code,400)


if __name__ == "__main__":
    unittest.main()
